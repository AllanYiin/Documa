"""Independent text/cell/bbox comparisons for the user-authorized Office corpus."""
from __future__ import annotations
import argparse
import os
from collections import Counter
from datetime import date, datetime
import json
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[1]
CORPUS = Path(os.environ.get("DOCUMA_OFFICE_CORPUS_DIR", str(ROOT / ".local-corpus"))).resolve()


def normalized(value):
    return re.sub(r"\s+", "", str(value))


def verify(case):
    import rust_office
    path = CORPUS / case["path"]
    payload = list(rust_office.open(path, extract_images=False))
    units = [e["unit"] for e in payload if e["event"] == "unit"]
    failures, checked, geometry = [], 0, 0
    if path.suffix == ".docx":
        from docx import Document
        source = Document(path)
        expected = [p.text for p in source.paragraphs if p.text.strip()]

        def table_text(table):
            seen, out = set(), []
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    text = "\n".join(p.text for p in cell.paragraphs).strip()
                    if text:
                        out.append(text)
                    for nested in cell.tables:
                        out.extend(table_text(nested))
            return out

        for table in source.tables:
            expected.extend(table_text(table))
        actual = [b["text"] for u in units for b in u["blocks"] if b["kind"] != "table" and any("word/document.xml" in r for r in b["source_refs"])]
        actual.extend(c["text"] for u in units for t in u["tables"] for c in t["metadata"]["cells"] if c["text"])
        if Counter(map(normalized, expected)) != Counter(map(normalized, actual)):
            failures.append("DOCX primary text differs from python-docx")
        checked = len(expected)
        reference = "python-docx"
    elif path.suffix == ".pptx":
        from pptx import Presentation
        source = Presentation(path)
        if len(source.slides) != len(units):
            failures.append("Slide count mismatch")
        for slide, unit in zip(source.slides, units):
            expected = []

            def leaves(shapes):
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        yield from leaves(shape.shapes)
                    else:
                        yield shape

            for index, shape in enumerate(leaves(slide.shapes)):
                if shape.has_text_frame and shape.text.strip():
                    expected.append(shape.text)
                    block = next((b for b in unit["blocks"] if b["metadata"].get("z_order") == index + 1), None)
                    bbox = [shape.left / 12700, shape.top / 12700, (shape.left + shape.width) / 12700, (shape.top + shape.height) / 12700]
                    if not block or not block["bbox"] or any(abs(a-b) > 0.01 for a, b in zip(block["bbox"], bbox)):
                        failures.append(f"Slide {unit['number']} shape {index+1} bbox mismatch")
                    geometry += 1
                elif shape.has_table:
                    expected.extend(c.text for row in shape.table.rows for c in row.cells if not c.is_spanned and c.text.strip())
            actual = [b["text"] for b in unit["blocks"] if b["kind"] != "table" and b["text"] and b["metadata"].get("source_type") != "speaker_notes"]
            actual.extend(c["text"] for t in unit["tables"] for c in t["metadata"]["cells"] if c["text"])
            if Counter(map(normalized, expected)) != Counter(map(normalized, actual)):
                failures.append(f"Slide {unit['number']} text differs from python-pptx")
            checked += len(expected)
        reference = "python-pptx"
    else:
        if path.suffix == ".xlsx":
            import openpyxl
            book = openpyxl.load_workbook(path, data_only=True)
            reference = "openpyxl"
        else:
            import xlrd
            book = xlrd.open_workbook(path, formatting_info=True)
            reference = "xlrd"
        for unit in units:
            records = {(c["row"]-1, c["column"]-1): c for b in unit["blocks"] for c in b["metadata"].get("cell_records", [])}
            if path.suffix == ".xlsx":
                sheet = book[unit["label"]]
                values = {(c.row-1, c.column-1): c.value for row in sheet for c in row if c.value is not None}
            else:
                sheet = book.sheet_by_name(unit["label"])
                values = {(r, c): sheet.cell_value(r, c) for r in range(sheet.nrows) for c in range(sheet.ncols) if sheet.cell_type(r, c) not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)}
            for pos, value in values.items():
                record = records.get(pos)
                correct = record is not None and record["raw_value"] == value
                if record is not None and isinstance(value, (date, datetime)):
                    correct = (record.get("date_iso") or "").startswith(value.isoformat())
                if not correct:
                    failures.append(f"{unit['label']} {pos}: raw value mismatch")
                checked += 1
            extras = set(records) - set(values)
            if any(records[pos]["raw_value"] is not None for pos in extras):
                failures.append(f"{unit['label']}: extra nonempty cells")
    return {"file": case["path"], "format": case["format"], "reference": reference, "items_checked": checked, "bbox_checked": geometry, "status": "PASS" if not failures else "FAIL", "failures": failures}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--require-real-world", action="store_true")
    parser.add_argument("--out", type=Path, default=ROOT / "reports" / "real-world-parity.json")
    args = parser.parse_args()
    manifest_path = CORPUS / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.is_file() else {"fixtures": []}
    results = [verify(case) for case in manifest["fixtures"]]
    failed = any(r["status"] != "PASS" for r in results)
    status = "FAIL" if failed or (not results and args.require_real_world) else "PASS" if results else "SKIPPED"
    report = {"status": status, "cases": results, "limitations": ["No private corpus is distributed with this repository.", "Text comparisons normalize whitespace and do not prove visual fidelity."]}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return int(status == "FAIL")


if __name__ == "__main__":
    raise SystemExit(main())
