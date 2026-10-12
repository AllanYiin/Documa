"""Generate small deterministic Office fixtures used by Rust and Python smoke tests."""

from __future__ import annotations

import base64
import hashlib
from datetime import datetime
import json
import re
from pathlib import Path
import zipfile


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures"
FIXED_TIMESTAMP = datetime(2000, 1, 1)
TINY_IMAGE = FIXTURES / ".fixture-image.png"
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def normalize_ooxml(path: Path) -> None:
    """Rewrite an OOXML ZIP with stable entry order, timestamps, and attributes."""

    normalized = path.with_suffix(path.suffix + ".normalized")
    with zipfile.ZipFile(path, "r") as source:
        entries = [(name, source.read(name)) for name in sorted(source.namelist())]
    canonical_entries = []
    for name, data in entries:
        if name == "docProps/core.xml":
            data = re.sub(
                rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                rb"\g<1>2000-01-01T00:00:00Z\g<3>",
                data,
            )
        canonical_entries.append((name, data))
    with zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name, data in canonical_entries:
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, data)
    normalized.replace(path)


def make_docx() -> None:
    from docx import Document

    document = Document()
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    document.add_heading("策略總覽", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Rust Office ").bold = True
    paragraph.add_run("保留文件結構與繁體中文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "格式"
    table.cell(0, 1).text = "支援"
    table.cell(1, 0).text = "DOCX"
    table.cell(1, 1).text = "段落與表格"
    document.save(FIXTURES / "smoke.docx")
    normalize_ooxml(FIXTURES / "smoke.docx")


def make_xlsx() -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    workbook = Workbook()
    workbook.properties.created = FIXED_TIMESTAMP
    workbook.properties.modified = FIXED_TIMESTAMP
    sheet = workbook.active
    sheet.title = "風險清單"
    sheet.append(["項目", "分數", "加權"])
    sheet.append(["流動性", 42, "=B2*2"])
    sheet["A2"].comment = Comment("重要風險", "Documa")
    sheet.merge_cells("A4:C4")
    sheet["A4"] = "合併儲存格"
    hidden = workbook.create_sheet("隱藏資料")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "不應預設顯示"
    workbook.save(FIXTURES / "smoke.xlsx")
    normalize_ooxml(FIXTURES / "smoke.xlsx")


def make_xls() -> None:
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("BIFF8資料")
    values = [
        ("項目", "分數", "加權"),
        ("相容性", 21, xlwt.Formula("B2*2")),
    ]
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            sheet.write(row_index, column_index, value)
    sheet.write_merge(3, 3, 0, 2, "BIFF8 合併儲存格")
    workbook.save(str(FIXTURES / "smoke.xls"))


def make_pptx() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.core_properties.created = FIXED_TIMESTAMP
    presentation.core_properties.modified = FIXED_TIMESTAMP
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "產品路線"
    slide.placeholders[1].text = "支援 Rust-first Office parsing"
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(4), Inches(5), Inches(1.2)
    ).table
    table.cell(0, 0).text = "格式"
    table.cell(0, 1).text = "狀態"
    table.cell(1, 0).text = "PPTX"
    table.cell(1, 1).text = "已解析"
    slide.notes_slide.notes_text_frame.text = "講者備註：下一階段接入 Documa。"
    presentation.save(FIXTURES / "smoke.pptx")
    normalize_ooxml(FIXTURES / "smoke.pptx")


def add_external_hyperlink(paragraph, text: str, target: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship_id = paragraph.part.relate_to(
        target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_docx_variants() -> None:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    document.add_heading("Unicode 與清單", level=1)
    document.add_paragraph("第一層項目：風險 αβγ", style="List Bullet")
    document.add_paragraph("第二層項目：控制措施 ✓", style="List Bullet 2")
    document.add_paragraph("Emoji 與多語：臺灣／日本語／한국어／🚀")
    path = FIXTURES / "docx-unicode-list.docx"
    document.save(path)
    normalize_ooxml(path)

    document = Document()
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    document.add_heading("合併表格", level=1)
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "跨欄表頭"
    for column, value in enumerate(("項目", "基準", "結果")):
        table.cell(1, column).text = value
    for column, value in enumerate(("資本", "100", "120")):
        table.cell(2, column).text = value
    path = FIXTURES / "docx-table-merged.docx"
    document.save(path)
    normalize_ooxml(path)

    document = Document()
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    section = document.sections[0]
    section.header.paragraphs[0].text = "頁首證據：機密"
    section.footer.paragraphs[0].text = "頁尾證據：第 N 頁"
    document.add_heading("章節與頁首頁尾", level=1)
    document.add_paragraph("正文保留 logical flow，不虛構頁碼。")
    path = FIXTURES / "docx-header-footer.docx"
    document.save(path)
    normalize_ooxml(path)

    document = Document()
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    document.add_heading("圖片與外部連結", level=1)
    document.add_picture(str(TINY_IMAGE), width=Inches(0.25))
    paragraph = document.add_paragraph("資產後方包含")
    add_external_hyperlink(paragraph, "外部參考", "https://example.invalid/rust-office")
    path = FIXTURES / "docx-image-link.docx"
    document.save(path)
    normalize_ooxml(path)

    (FIXTURES / "docx-corrupt.docx").write_bytes(b"PK\x03\x04truncated-docx")


def make_xls_variants() -> None:
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Unicode")
    sheet.write(0, 0, "供應鏈")
    sheet.write(1, 0, "臺灣／日本語／한국어")
    workbook.save(str(FIXTURES / "xls-unicode.xls"))

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("公式")
    sheet.write(0, 0, "折扣")
    sheet.write(0, 1, "原價")
    sheet.write(0, 2, "結果")
    sheet.write(1, 0, 0.8)
    sheet.write(1, 1, 100)
    sheet.write(1, 2, xlwt.Formula("A2*B2"))
    workbook.save(str(FIXTURES / "xls-formula.xls"))

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("合併")
    sheet.write_merge(0, 0, 0, 3, "BIFF8 合併標題")
    sheet.write(1, 0, "欄位")
    sheet.write(1, 1, "內容")
    workbook.save(str(FIXTURES / "xls-merged.xls"))

    workbook = xlwt.Workbook()
    visible = workbook.add_sheet("公開")
    visible.write(0, 0, "公開資料")
    hidden = workbook.add_sheet("隱藏")
    hidden.write(0, 0, "隱藏資料")
    hidden.visibility = 1
    workbook.save(str(FIXTURES / "xls-multisheet.xls"))

    (FIXTURES / "xls-invalid.xls").write_bytes(b"not-an-ole-biff8-workbook")


def save_xlsx(workbook, name: str) -> Path:
    workbook.properties.created = FIXED_TIMESTAMP
    workbook.properties.modified = FIXED_TIMESTAMP
    path = FIXTURES / name
    workbook.save(path)
    normalize_ooxml(path)
    return path


def make_xlsx_variants() -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Unicode公式"
    sheet.append(["項目", "單價", "數量", "總額"])
    sheet.append(["臺灣市場", 25, 4, "=B2*C2"])
    sheet.append(["日本語", 10, 3, "=B3*C3"])
    save_xlsx(workbook, "xlsx-unicode-formula.xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合併連結"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "合併範圍"
    sheet["A2"] = "外部資料"
    sheet["A2"].hyperlink = "https://example.invalid/sheet"
    sheet["A2"].comment = Comment("只保留 metadata", "Documa")
    save_xlsx(workbook, "xlsx-merged-link.xlsx")

    workbook = Workbook()
    visible = workbook.active
    visible.title = "可見"
    visible["A1"] = "可見工作表"
    hidden = workbook.create_sheet("隱藏")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "隱藏工作表"
    save_xlsx(workbook, "xlsx-hidden.xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "圖表清單"
    sheet.append(["季度", "營收"])
    sheet.append(["Q1", 10])
    sheet.append(["Q2", 14])
    table = Table(displayName="QuarterlyRevenue", ref="A1:B3")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    chart = BarChart()
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D2")
    base = save_xlsx(workbook, "xlsx-table-chart.xlsx")
    add_unsafe_zip_entry(base, FIXTURES / "xlsx-path-traversal.xlsx")


def save_pptx(presentation, name: str) -> Path:
    presentation.core_properties.created = FIXED_TIMESTAMP
    presentation.core_properties.modified = FIXED_TIMESTAMP
    path = FIXTURES / name
    presentation.save(path)
    normalize_ooxml(path)
    return path


def make_pptx_variants() -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = Presentation()
    for title, body in (
        ("第一張投影片", "摘要與範圍"),
        ("第二張投影片", "決策與後續行動"),
        ("第三張投影片", "附錄"),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    save_pptx(presentation, "pptx-multislide.pptx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Unicode 與項目符號"
    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()
    text_frame.paragraphs[0].text = "第一層項目：臺灣"
    nested = text_frame.add_paragraph()
    nested.text = "巢狀項目：日本語／한국어／🚀"
    nested.level = 1
    save_pptx(presentation, "pptx-unicode-bullets.pptx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "會議結論"
    table = slide.shapes.add_table(
        3, 2, Inches(1), Inches(1.8), Inches(6), Inches(2)
    ).table
    for row, values in enumerate(
        (("項目", "狀態"), ("Parser", "完成"), ("Gate", "待驗證"))
    ):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    slide.notes_slide.notes_text_frame.text = "講者備註：確認 release gate。"
    save_pptx(presentation, "pptx-table-notes.pptx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "圖片與連結"
    slide.shapes.add_picture(str(TINY_IMAGE), Inches(1), Inches(1.5), width=Inches(0.5))
    link_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(2), Inches(1.5), Inches(2), Inches(0.8)
    )
    link_shape.text = "外部參考"
    link_shape.click_action.hyperlink.address = "https://example.invalid/slide"
    base = save_pptx(presentation, "pptx-image-link.pptx")
    add_unsafe_zip_entry(base, FIXTURES / "pptx-path-traversal.pptx")


def add_unsafe_zip_entry(source: Path, destination: Path) -> None:
    with zipfile.ZipFile(source, "r") as archive:
        entries = [(name, archive.read(name)) for name in archive.namelist()]
    entries.append(("../escape.xml", b"<escape/>"))
    with zipfile.ZipFile(
        destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as archive:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(info, data)
    normalize_ooxml(destination)


FIXTURE_CASES = [
    ("smoke.docx", "docx", ["unicode", "heading", "runs", "table"], "繁體中文", None),
    (
        "docx-unicode-list.docx",
        "docx",
        ["unicode", "heading", "list"],
        "Unicode 與清單",
        None,
    ),
    ("docx-table-merged.docx", "docx", ["table", "merged_cells"], "跨欄表頭", None),
    (
        "docx-header-footer.docx",
        "docx",
        ["sections", "header", "footer"],
        "頁首證據",
        None,
    ),
    ("docx-image-link.docx", "docx", ["image", "external_link"], "外部參考", None),
    ("docx-corrupt.docx", "docx", ["corrupt", "security"], None, "CORRUPT_PACKAGE"),
    ("smoke.xls", "xls", ["unicode", "formula", "merged_cells"], "BIFF8", None),
    ("xls-unicode.xls", "xls", ["unicode"], "供應鏈", None),
    ("xls-formula.xls", "xls", ["formula", "cached_value"], "折扣", None),
    ("xls-merged.xls", "xls", ["merged_cells"], "BIFF8 合併標題", None),
    (
        "xls-multisheet.xls",
        "xls",
        ["multiple_sheets", "hidden_sheet"],
        "公開資料",
        None,
    ),
    ("xls-invalid.xls", "xls", ["corrupt", "format_mismatch"], None, "FORMAT_MISMATCH"),
    (
        "smoke.xlsx",
        "xlsx",
        ["unicode", "formula", "merged_cells", "comment", "hidden_sheet"],
        "流動性",
        None,
    ),
    (
        "xlsx-unicode-formula.xlsx",
        "xlsx",
        ["unicode", "formula", "cached_value"],
        "總額",
        None,
    ),
    (
        "xlsx-merged-link.xlsx",
        "xlsx",
        ["merged_cells", "hyperlink", "comment"],
        "外部資料",
        None,
    ),
    ("xlsx-hidden.xlsx", "xlsx", ["hidden_sheet"], "可見工作表", None),
    (
        "xlsx-table-chart.xlsx",
        "xlsx",
        ["structured_table", "chart_inventory"],
        "季度",
        None,
    ),
    (
        "xlsx-path-traversal.xlsx",
        "xlsx",
        ["security", "path_traversal"],
        None,
        "ZIP_PATH_TRAVERSAL",
    ),
    (
        "smoke.pptx",
        "pptx",
        ["unicode", "placeholder_geometry", "table", "speaker_notes"],
        "產品路線",
        None,
    ),
    (
        "pptx-multislide.pptx",
        "pptx",
        ["multiple_slides", "reading_order"],
        "第二張投影片",
        None,
    ),
    ("pptx-unicode-bullets.pptx", "pptx", ["unicode", "bullets"], "巢狀項目", None),
    ("pptx-table-notes.pptx", "pptx", ["table", "speaker_notes"], "會議結論", None),
    (
        "pptx-image-link.pptx",
        "pptx",
        ["image", "external_link", "bbox"],
        "外部參考",
        None,
    ),
    (
        "pptx-path-traversal.pptx",
        "pptx",
        ["security", "path_traversal"],
        None,
        "ZIP_PATH_TRAVERSAL",
    ),
]


def write_manifest() -> None:
    fixtures = []
    for name, format_name, coverage, needle, expected_error in FIXTURE_CASES:
        path = FIXTURES / name
        entry = {
            "path": name,
            "format": format_name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "coverage": coverage,
            "provenance": "project-generated deterministic synthetic fixture",
            "license": "MIT",
        }
        if format_name == "xls" and expected_error is None:
            entry["variant"] = "BIFF8"
        if needle is not None:
            entry["expected_needle"] = needle
        if expected_error is not None:
            entry["expected_error"] = expected_error
        fixtures.append(entry)
    manifest = {
        "schema_version": 1,
        "generator": "scripts/generate_fixtures.py",
        "deterministic": True,
        "quality_gate": {
            "required_committed_fixtures": 24,
            "current_committed_fixtures": len(fixtures),
            "status": "complete" if len(fixtures) >= 24 else "partial",
        },
        "fixtures": fixtures,
    }
    (FIXTURES / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    TINY_IMAGE.write_bytes(TINY_PNG)
    try:
        make_docx()
        make_docx_variants()
        make_xls()
        make_xls_variants()
        make_xlsx()
        make_xlsx_variants()
        make_pptx()
        make_pptx_variants()
        write_manifest()
    finally:
        TINY_IMAGE.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
